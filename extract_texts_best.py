#!/usr/bin/env python3
"""Best-of text extractor with OCR, parallelism, and safe output naming.

Usage:
  python extract_texts_best.py /path/to/input /path/to/output
  python extract_texts_best.py /path/to/input /path/to/output --parallel
"""

from __future__ import annotations

import argparse
import json
import re
import sys
import tempfile
import zipfile
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable, Optional

try:
    from tqdm import tqdm
    HAS_TQDM = True
except ImportError:
    HAS_TQDM = False

    def tqdm(iterable, **kwargs):
        return iterable

TEXT_EXTENSIONS = {
    ".txt",
    ".md",
    ".markdown",
    ".csv",
    ".tsv",
    ".json",
    ".xml",
    ".html",
    ".htm",
    ".rtf",
    ".log",
}

IMAGE_EXTENSIONS = {
    ".png",
    ".jpg",
    ".jpeg",
    ".tiff",
    ".tif",
    ".bmp",
    ".gif",
    ".webp",
}


@dataclass
class ExtractedText:
    source_zip: str
    source_file: str
    text: str
    method: str = "unknown"
    char_count: int = 0
    confidence: float = 0.0


def iter_zip_files(root: Path) -> Iterable[Path]:
    for path in root.rglob("*.zip"):
        if path.is_file():
            yield path


def read_text_file(path: Path) -> Optional[str]:
    try:
        return path.read_text(encoding="utf-8")
    except UnicodeDecodeError:
        try:
            return path.read_text(encoding="latin-1", errors="replace")
        except Exception:
            return None


def preprocess_image_for_ocr(image):
    try:
        import cv2
        import numpy as np

        img_array = np.array(image)
        if len(img_array.shape) == 3:
            gray = cv2.cvtColor(img_array, cv2.COLOR_RGB2GRAY)
        else:
            gray = img_array

        clahe = cv2.createCLAHE(clipLimit=2.0, tileGridSize=(8, 8))
        enhanced = clahe.apply(gray)
        denoised = cv2.fastNlMeansDenoising(enhanced, None, 10, 7, 21)
        _, binary = cv2.threshold(
            denoised, 0, 255, cv2.THRESH_BINARY + cv2.THRESH_OTSU
        )
        return binary
    except ImportError:
        return image
    except Exception:
        return image


def extract_image_text(path: Path, ocr_lang: str) -> Optional[tuple[str, float]]:
    try:
        from PIL import Image
    except ImportError:
        return None

    try:
        import pytesseract
    except ImportError:
        return None

    try:
        with Image.open(path) as image:
            if image.mode not in ("RGB", "L"):
                image = image.convert("RGB")

            results = []

            try:
                text1 = pytesseract.image_to_string(image, lang=ocr_lang).strip()
                if text1:
                    try:
                        data = pytesseract.image_to_data(
                            image, lang=ocr_lang, output_type=pytesseract.Output.DICT
                        )
                        confidences = [
                            int(conf) for conf in data["conf"] if conf != "-1"
                        ]
                        avg_conf = sum(confidences) / len(confidences) if confidences else 0
                        results.append((text1, avg_conf, "direct"))
                    except Exception:
                        results.append((text1, 75.0, "direct"))
            except Exception as exc:
                print(f"  [Warning] Direct OCR failed: {exc}")

            try:
                preprocessed = preprocess_image_for_ocr(image)
                text2 = pytesseract.image_to_string(
                    preprocessed, lang=ocr_lang
                ).strip()
                if text2:
                    results.append((text2, 80.0, "preprocessed"))
            except Exception as exc:
                print(f"  [Warning] Preprocessed OCR failed: {exc}")

            try:
                if image.width < 1000 or image.height < 1000:
                    scale_factor = max(1000 / image.width, 1000 / image.height)
                    new_size = (
                        int(image.width * scale_factor),
                        int(image.height * scale_factor),
                    )
                    upscaled = image.resize(new_size, Image.Resampling.LANCZOS)
                    text3 = pytesseract.image_to_string(
                        upscaled, lang=ocr_lang
                    ).strip()
                    if text3:
                        results.append((text3, 70.0, "upscaled"))
            except Exception as exc:
                print(f"  [Warning] Upscaled OCR failed: {exc}")

            if results:
                best = max(results, key=lambda x: len(x[0]) * (x[1] / 100))
                return (best[0], best[1])
            return None
    except Exception as exc:
        print(f"  [Error] Image OCR failed for {path}: {exc}")
        return None


def extract_pdf_text(path: Path, ocr_lang: str) -> Optional[tuple[str, str]]:
    try:
        import pdfplumber

        parts = []
        with pdfplumber.open(str(path)) as pdf:
            for page in pdf.pages:
                text = page.extract_text() or ""
                parts.append(text)
        text = "\n".join(parts).strip()
        if text and len(text) > 20:
            return (text, "pdfplumber")
    except ImportError:
        pass
    except Exception as exc:
        print(f"  [Warning] pdfplumber failed: {exc}")

    try:
        import PyPDF2

        reader = PyPDF2.PdfReader(str(path))
        parts = []
        for page in reader.pages:
            parts.append(page.extract_text() or "")
        text = "\n".join(parts).strip()
        if text and len(text) > 20:
            return (text, "PyPDF2")
    except ImportError:
        pass
    except Exception as exc:
        print(f"  [Warning] PyPDF2 failed: {exc}")

    try:
        from pdfminer.high_level import extract_text

        text = extract_text(str(path)).strip()
        if text and len(text) > 20:
            return (text, "pdfminer")
    except ImportError:
        pass
    except Exception as exc:
        print(f"  [Warning] pdfminer failed: {exc}")

    try:
        import pdf2image
        import pytesseract

        print("  [Info] Attempting OCR on PDF (may be slow)...")
        images = pdf2image.convert_from_path(str(path), dpi=300)
        parts = []
        for i, image in enumerate(images):
            try:
                text = pytesseract.image_to_string(image, lang=ocr_lang).strip()
                if text:
                    parts.append(text)
            except Exception as exc:
                print(f"  [Warning] OCR failed on page {i+1}: {exc}")
        text = "\n\n".join(parts).strip()
        if text:
            return (text, "pdf_ocr")
    except ImportError:
        pass
    except Exception as exc:
        print(f"  [Warning] PDF OCR failed: {exc}")

    return None


def extract_docx_text(path: Path) -> Optional[str]:
    try:
        import docx

        document = docx.Document(str(path))
        parts = []
        for paragraph in document.paragraphs:
            if paragraph.text.strip():
                parts.append(paragraph.text)
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        parts.append(cell.text)
        return "\n".join(parts).strip()
    except ImportError:
        pass
    except Exception:
        pass

    try:
        with zipfile.ZipFile(path) as docx_zip:
            xml_bytes = docx_zip.read("word/document.xml")
        from xml.etree import ElementTree

        root = ElementTree.fromstring(xml_bytes)
        texts = []
        for node in root.iter():
            if node.tag.endswith("}t") and node.text:
                texts.append(node.text)
        return " ".join(texts).strip()
    except Exception:
        return None


def sanitize_filename(name: str) -> str:
    sanitized = re.sub(r"[^A-Za-z0-9._-]+", "_", name)
    return sanitized.strip("_") or "document"


def unique_output_path(directory: Path, base_name: str) -> Path:
    candidate = directory / base_name
    if not candidate.exists():
        return candidate
    stem = candidate.stem
    suffix = candidate.suffix
    for idx in range(1, 10000):
        candidate = directory / f"{stem}_{idx}{suffix}"
        if not candidate.exists():
            return candidate
    raise RuntimeError(f"Unable to find unique filename for {base_name}")


def extract_text_from_file(path: Path, ocr_lang: str) -> Optional[ExtractedText]:
    ext = path.suffix.lower()
    text = None
    method = "unknown"
    confidence = 0.0

    if ext in TEXT_EXTENSIONS:
        text = read_text_file(path)
        method = "plaintext"
        confidence = 100.0
    elif ext == ".pdf":
        result = extract_pdf_text(path, ocr_lang)
        if result:
            text, method = result
            confidence = 95.0 if method != "pdf_ocr" else 75.0
    elif ext == ".docx":
        text = extract_docx_text(path)
        method = "docx"
        confidence = 95.0
    elif ext in IMAGE_EXTENSIONS:
        result = extract_image_text(path, ocr_lang)
        if result:
            text, confidence = result
            method = "ocr"

    if text:
        return ExtractedText(
            source_zip="",
            source_file=str(path),
            text=text,
            method=method,
            char_count=len(text),
            confidence=confidence,
        )
    return None


def process_zip(
    zip_path: Path,
    output_dir: Path,
    ocr_lang: str,
    verbose: bool = True,
) -> list[ExtractedText]:
    extracted: list[ExtractedText] = []

    if verbose:
        print(f"\n📦 Processing: {zip_path.name}")

    with tempfile.TemporaryDirectory() as tmp_dir:
        try:
            with zipfile.ZipFile(zip_path) as archive:
                archive.extractall(tmp_dir)
        except zipfile.BadZipFile:
            print(f"  ❌ Invalid zip file: {zip_path}")
            return []

        files = list(Path(tmp_dir).rglob("*"))
        if verbose:
            print(f"  Found {len([f for f in files if f.is_file()])} files")

        iterator = (
            tqdm(files, desc="  Extracting", leave=False)
            if HAS_TQDM and verbose
            else files
        )

        for file_path in iterator:
            if not file_path.is_file():
                continue

            result = extract_text_from_file(file_path, ocr_lang)
            if result is None:
                continue

            result.source_zip = str(zip_path)
            result.source_file = str(file_path.relative_to(tmp_dir))
            extracted.append(result)

            if verbose and not HAS_TQDM:
                print(
                    f"    ✓ {result.source_file} ({result.char_count} chars, {result.method})"
                )

    if extracted:
        zip_output_dir = output_dir / sanitize_filename(zip_path.stem)
        zip_output_dir.mkdir(parents=True, exist_ok=True)
        for item in extracted:
            output_name = sanitize_filename(Path(item.source_file).stem) + ".txt"
            output_path = unique_output_path(zip_output_dir, output_name)
            output_path.write_text(item.text, encoding="utf-8")

    if verbose:
        print(f"  ✅ Extracted {len(extracted)} documents")

    return extracted


def write_jsonl(records: Iterable[ExtractedText], output_path: Path) -> None:
    with output_path.open("w", encoding="utf-8") as handle:
        for record in records:
            handle.write(
                json.dumps(
                    {
                        "source_zip": record.source_zip,
                        "source_file": record.source_file,
                        "text": record.text,
                        "method": record.method,
                        "char_count": record.char_count,
                        "confidence": record.confidence,
                    },
                    ensure_ascii=False,
                )
                + "\n"
            )


def generate_summary_report(records: list[ExtractedText], output_dir: Path) -> None:
    from collections import Counter

    total_docs = len(records)
    total_chars = sum(r.char_count for r in records)
    methods = Counter(r.method for r in records)
    avg_confidence = (
        sum(r.confidence for r in records) / total_docs if total_docs > 0 else 0
    )

    report = f"""
# Extraction Summary Report

## Overall Statistics
- Total documents processed: {total_docs}
- Total characters extracted: {total_chars:,}
- Average confidence: {avg_confidence:.1f}%

## Extraction Methods
"""
    for method, count in methods.most_common():
        percentage = (count / total_docs * 100) if total_docs > 0 else 0
        report += f"- {method}: {count} documents ({percentage:.1f}%)\n"

    report += "\n## Documents by Confidence\n"
    high_conf = sum(1 for r in records if r.confidence >= 90)
    med_conf = sum(1 for r in records if 70 <= r.confidence < 90)
    low_conf = sum(1 for r in records if r.confidence < 70)

    report += f"- High (≥90%): {high_conf}\n"
    report += f"- Medium (70-89%): {med_conf}\n"
    report += f"- Low (<70%): {low_conf}\n"

    report += "\n## Documents Requiring Review (low confidence or OCR)\n"
    review_docs = [r for r in records if r.confidence < 80 or "ocr" in r.method]
    for doc in review_docs[:20]:
        report += (
            f"- {doc.source_file} ({doc.method}, {doc.confidence:.0f}% conf, "
            f"{doc.char_count} chars)\n"
        )

    if len(review_docs) > 20:
        report += f"... and {len(review_docs) - 20} more\n"

    report_path = output_dir / "extraction_report.md"
    report_path.write_text(report, encoding="utf-8")
    print(f"\n📊 Summary report: {report_path}")


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Best text extractor with advanced OCR and safe outputs."
    )
    parser.add_argument("input_path", help="Path containing zip files")
    parser.add_argument(
        "output_dir",
        nargs="?",
        default="extracted_text",
        help="Directory to write extracted text",
    )
    parser.add_argument(
        "--parallel",
        action="store_true",
        help="Process zip files in parallel (faster but uses more memory)",
    )
    parser.add_argument(
        "--workers",
        type=int,
        default=4,
        help="Number of parallel workers (default: 4)",
    )
    parser.add_argument(
        "--quiet",
        action="store_true",
        help="Minimal output",
    )
    parser.add_argument(
        "--ocr-lang",
        default="eng+fra",
        help="OCR languages for Tesseract (default: eng+fra)",
    )
    args = parser.parse_args()

    input_path = Path(args.input_path).expanduser().resolve()
    output_dir = Path(args.output_dir).expanduser().resolve()

    if not input_path.exists():
        print(f"❌ Input path not found: {input_path}", file=sys.stderr)
        return 1

    output_dir.mkdir(parents=True, exist_ok=True)

    print("🚀 Starting text extraction...")
    print(f"📂 Input: {input_path}")
    print(f"📂 Output: {output_dir}")

    zip_files = list(iter_zip_files(input_path))
    if not zip_files:
        print("❌ No zip files found", file=sys.stderr)
        return 1

    print(f"📦 Found {len(zip_files)} zip file(s)")

    all_records: list[ExtractedText] = []

    if args.parallel:
        print(f"⚡ Using {args.workers} parallel workers")
        with ThreadPoolExecutor(max_workers=args.workers) as executor:
            futures = {
                executor.submit(process_zip, zp, output_dir, args.ocr_lang, not args.quiet): zp
                for zp in zip_files
            }

            for future in (
                tqdm(as_completed(futures), total=len(futures), desc="Processing zips")
                if HAS_TQDM
                else as_completed(futures)
            ):
                try:
                    records = future.result()
                    all_records.extend(records)
                except Exception as exc:
                    zip_path = futures[future]
                    print(f"❌ Error processing {zip_path}: {exc}", file=sys.stderr)
    else:
        for zip_path in zip_files:
            try:
                records = process_zip(zip_path, output_dir, args.ocr_lang, not args.quiet)
                all_records.extend(records)
            except Exception as exc:
                print(f"❌ Error processing {zip_path}: {exc}", file=sys.stderr)

    jsonl_path = output_dir / "extracted_text.jsonl"
    write_jsonl(all_records, jsonl_path)

    if not args.quiet:
        generate_summary_report(all_records, output_dir)

    print(f"\n✅ Done! Processed {len(all_records)} documents")
    print(f"📄 JSONL output: {jsonl_path}")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
